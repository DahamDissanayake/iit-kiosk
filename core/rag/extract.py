"""Turns uploaded files and pasted text into (heading, text) sections."""

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from markdown_it import MarkdownIt
from pypdf import PdfReader


@dataclass(frozen=True)
class Section:
    heading: str
    text: str


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _group(items: list[tuple[str, str]]) -> list[Section]:
    """items are ("h", heading) or ("p", paragraph); paragraphs join under the last heading."""
    sections: list[Section] = []
    heading, body = "", []
    for kind, value in items:
        if kind == "h":
            if body:
                sections.append(Section(heading, _norm(" ".join(body))))
            heading, body = _norm(value), []
        elif value.strip():
            body.append(value)
    if body:
        sections.append(Section(heading, _norm(" ".join(body))))
    return sections


def sections_from_markdown(text: str) -> list[Section]:
    tokens = MarkdownIt().parse(text)
    items: list[tuple[str, str]] = []
    in_heading = False
    for tok in tokens:
        if tok.type == "heading_open":
            in_heading = True
        elif tok.type == "heading_close":
            in_heading = False
        elif tok.type == "inline":
            plain = "".join(
                c.content for c in (tok.children or []) if c.type in ("text", "code_inline")
            )
            items.append(("h" if in_heading else "p", plain))
    return _group(items)


def sections_from_text(text: str) -> list[Section]:
    return _group([("p", p) for p in re.split(r"\n\s*\n", text)])


def sections_from_docx(path: Path) -> list[Section]:
    doc = Document(str(path))
    items: list[tuple[str, str]] = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        kind = "h" if style.startswith(("Heading", "Title")) else "p"
        items.append((kind, p.text))
    for table in doc.tables:
        for row in table.rows:
            items.append(("p", " | ".join(cell.text for cell in row.cells)))
    return _group(items)


def sections_from_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    return [
        Section("", _norm(text))
        for page in reader.pages
        if (text := page.extract_text() or "").strip()
    ]


def extract(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        sections = sections_from_pdf(path)
    elif suffix == ".docx":
        sections = sections_from_docx(path)
    elif suffix in (".md", ".markdown"):
        sections = sections_from_markdown(path.read_text(encoding="utf-8"))
    elif suffix == ".txt":
        sections = sections_from_text(path.read_text(encoding="utf-8"))
    else:
        raise ValueError(f"Unsupported file type: {suffix or 'none'}. Use PDF, DOCX, MD or TXT.")
    if not sections:
        raise ValueError("The file has no text that can be read (is it a scanned image?).")
    return sections
