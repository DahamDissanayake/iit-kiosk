import pytest
from docx import Document

from core.rag.chunk import Chunk, chunk_sections, count_tokens, qa_chunks
from core.rag.extract import Section, extract, sections_from_markdown


def minimal_pdf(text: str | None) -> bytes:
    """A one-page PDF; with text=None the page has no text at all."""
    stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out


# ---- extraction ----


def test_markdown_keeps_headings_with_their_text():
    md = "# Programmes\nWe teach **computing**.\n\n## Business\nAnd business too.\n"
    assert sections_from_markdown(md) == [
        Section("Programmes", "We teach computing."),
        Section("Business", "And business too."),
    ]


def test_markdown_without_headings_is_one_section():
    assert sections_from_markdown("Just a paragraph.\n\nAnother one.") == [
        Section("", "Just a paragraph. Another one.")
    ]


def test_docx_headings_and_paragraphs(tmp_path):
    doc = Document()
    doc.add_heading("Campus", level=1)
    doc.add_paragraph("The campus is in Colombo.")
    doc.add_heading("Library", level=2)
    doc.add_paragraph("Open until 8 pm.")
    path = tmp_path / "info.docx"
    doc.save(path)
    assert extract(path) == [
        Section("Campus", "The campus is in Colombo."),
        Section("Library", "Open until 8 pm."),
    ]


def test_pdf_text_is_extracted(tmp_path):
    path = tmp_path / "info.pdf"
    path.write_bytes(minimal_pdf("Hello campus"))
    sections = extract(path)
    assert "Hello campus" in sections[0].text


def test_pdf_without_text_is_rejected(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(minimal_pdf(None))
    with pytest.raises(ValueError, match="no text"):
        extract(path)


def test_unsupported_file_type_is_rejected(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError, match="Unsupported"):
        extract(path)


# ---- chunking ----


def test_short_section_is_one_chunk_with_heading():
    chunks = chunk_sections([Section("Campus", "The campus is in Colombo.")])
    assert chunks == [Chunk("Campus", "Campus\nThe campus is in Colombo.", 7)]


def test_long_section_is_split_with_overlap_and_heading_in_every_chunk():
    words = [f"w{i}" for i in range(600)]
    chunks = chunk_sections([Section("Research", " ".join(words))], size=250, overlap=40)
    assert len(chunks) >= 3
    for c in chunks:
        assert c.text.startswith("Research\n")
        assert c.tokens <= 250
    first = chunks[0].text.split("\n", 1)[1].split()
    second = chunks[1].text.split("\n", 1)[1].split()
    shared = [w for w in second if w in set(first)]
    assert 30 <= len(shared) <= 45
    last = chunks[-1].text.split()
    assert last[-1] == "w599"


def test_empty_sections_are_skipped():
    assert chunk_sections([Section("Empty", "   ")]) == []


def test_qa_pairs_are_one_chunk_each():
    chunks = qa_chunks([("Where is IIT?", "In Colombo."), ("Is there parking?", "Yes.")])
    assert [c.text for c in chunks] == [
        "Q: Where is IIT?\nA: In Colombo.",
        "Q: Is there parking?\nA: Yes.",
    ]


def test_count_tokens_counts_words_and_punctuation():
    assert count_tokens("Hello, world!") == 4
